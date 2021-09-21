import os

import docx
import openpyxl

book = openpyxl.load_workbook("data/名单.xlsx")
sheet = book["Sheet1"]
word = docx.Document("data/模板.docx")
for i in range(2, 5):
    name = sheet[f"A{i}"].value
    word.paragraphs[0].runs[0].text = name  # 需要修改的位置
    if not os.path.exists('temp'):
        os.mkdir('temp')
    word.save(f"temp/{name}.docx")
